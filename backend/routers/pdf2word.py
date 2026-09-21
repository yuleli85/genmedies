import asyncio
import re
import uuid
from pathlib import Path

import fitz
from docx import Document
from fastapi import APIRouter, BackgroundTasks, File, HTTPException, UploadFile
from pydantic import BaseModel

from services.llm_service import extract_text_from_image, repair_ocr_text
from services.task_store import TaskStatus, create_task, get_task, update_task

router = APIRouter(prefix="/api/pdf2word", tags=["pdf2word"])

PDF2DOCX_IMPORT_ERROR = None
try:
    from pdf2docx import Converter
except ImportError as e:
    Converter = None
    PDF2DOCX_IMPORT_ERROR = str(e)

PDF_DIR = Path("uploads/pdf2word")
PDF_DIR.mkdir(parents=True, exist_ok=True)
OUTPUT_DIR = Path("uploads/pdf2word_output")
OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
PAGE_IMAGE_DIR = Path("uploads/pdf2word_pages")
PAGE_IMAGE_DIR.mkdir(parents=True, exist_ok=True)

MIN_TEXT_LENGTH = 80
MIN_TEXT_DENSITY = 0.00003
MIN_IMAGE_COVERAGE = 0.35


@router.post("/upload")
async def upload_pdf(file: UploadFile = File(...)):
    filename = file.filename or ""
    if not filename.lower().endswith(".pdf"):
        raise HTTPException(400, "仅支持 .pdf 文件")

    content = await file.read()
    if not content:
        raise HTTPException(400, "PDF 文件为空")

    file_id = f"pdf_{uuid.uuid4().hex[:12]}"
    save_path = PDF_DIR / f"{file_id}.pdf"
    save_path.write_bytes(content)

    return {
        "file_id": file_id,
        "filename": filename,
    }


class ConvertRequest(BaseModel):
    file_id: str


def _page_image_path(file_id: str, page_num: int) -> Path:
    page_dir = PAGE_IMAGE_DIR / file_id
    page_dir.mkdir(parents=True, exist_ok=True)
    return page_dir / f"page_{page_num:03d}.png"


def _normalize_page_text(text: str) -> str:
    text = text.replace("\x00", " ")
    text = re.sub(r"\s+", " ", text)
    return text.strip()


def _needs_ocr(page, text: str) -> bool:
    normalized = _normalize_page_text(text)
    text_len = len(normalized)
    page_rect = page.rect
    page_area = max(page_rect.width * page_rect.height, 1)
    image_coverage = 0.0
    image_blocks = 0

    for block in page.get_text("dict").get("blocks", []):
        if block.get("type") != 1:
            continue
        image_blocks += 1
        bbox = block.get("bbox") or [0, 0, 0, 0]
        x0, y0, x1, y1 = bbox
        block_area = max((x1 - x0) * (y1 - y0), 0)
        image_coverage += block_area / page_area

    text_density = text_len / page_area
    has_large_image = image_coverage >= MIN_IMAGE_COVERAGE
    has_too_little_text = text_len < MIN_TEXT_LENGTH or text_density < MIN_TEXT_DENSITY

    return has_too_little_text and has_large_image or text_len == 0


def _extract_embedded_page_text(pdf_path: Path):
    doc = fitz.open(pdf_path)
    pages = []
    try:
        for i, page in enumerate(doc, start=1):
            raw_text = page.get_text("text")
            text = _normalize_page_text(raw_text)
            pages.append({
                "page": i,
                "text": text,
                "needs_ocr": _needs_ocr(page, raw_text),
            })
    finally:
        doc.close()
    return pages


def _render_page_to_image(pdf_path: Path, file_id: str, page_num: int) -> Path:
    doc = fitz.open(pdf_path)
    try:
        page = doc.load_page(page_num - 1)
        pix = page.get_pixmap(matrix=fitz.Matrix(2, 2), alpha=False)
        out = _page_image_path(file_id, page_num)
        pix.save(out)
        return out
    finally:
        doc.close()


def _build_text_docx(output_path: Path, page_texts: list[dict]):
    doc = Document()
    for idx, item in enumerate(page_texts, start=1):
        doc.add_heading(f"第 {item['page']} 页", level=2)
        text = item.get("text", "").strip()
        if text:
            for para in [p for p in text.splitlines() if p.strip()]:
                doc.add_paragraph(para)
        else:
            doc.add_paragraph("未识别到可读文字")
        if idx < len(page_texts):
            doc.add_page_break()
    doc.save(output_path)


async def _ocr_pages(pdf_path: Path, file_id: str, pages: list[dict], task_id: str):
    total = len(pages)
    ocr_pages = []
    results = []
    for idx, item in enumerate(pages, start=1):
        if not item["needs_ocr"]:
            results.append({"page": item["page"], "text": item["text"], "source": "embedded"})
            continue

        update_task(
            task_id,
            status=TaskStatus.RUNNING,
            progress=30 + int(idx / max(total, 1) * 25),
            message=f"正在识别第 {item['page']} 页图片文字...",
        )
        image_path = _render_page_to_image(pdf_path, file_id, item["page"])
        text = await extract_text_from_image(str(image_path))

        update_task(
            task_id,
            status=TaskStatus.RUNNING,
            progress=55 + int(idx / max(total, 1) * 25),
            message=f"正在 AI 修复第 {item['page']} 页文字...",
        )
        repaired = await repair_ocr_text(text, str(image_path)) if text.strip() else text
        final_text = repaired.strip() or item["text"]

        ocr_pages.append(item["page"])
        results.append({"page": item["page"], "text": final_text, "source": "ocr"})
    return results, ocr_pages


def _convert_direct_pdf_to_docx(source_path: Path, output_path: Path):
    if Converter is None:
        raise RuntimeError(f"缺少 pdf2docx 依赖 ({PDF2DOCX_IMPORT_ERROR})")

    converter = Converter(str(source_path))
    try:
        converter.convert(str(output_path))
    finally:
        converter.close()


@router.post("/convert")
async def convert_pdf(req: ConvertRequest, bg: BackgroundTasks):
    source_path = PDF_DIR / f"{req.file_id}.pdf"
    if not source_path.exists():
        raise HTTPException(404, "PDF 文件不存在")

    task = create_task()

    async def _run():
        output_path = OUTPUT_DIR / f"{req.file_id}.docx"
        try:
            update_task(task.task_id, status=TaskStatus.RUNNING, progress=5, message="开始分析 PDF")
            pages = _extract_embedded_page_text(source_path)
            page_count = len(pages)
            needs_ocr = any(item["needs_ocr"] for item in pages)

            if not needs_ocr:
                update_task(task.task_id, progress=40, message="检测到文本型 PDF，直接转换")
                await asyncio.to_thread(_convert_direct_pdf_to_docx, source_path, output_path)
                update_task(
                    task.task_id,
                    status=TaskStatus.DONE,
                    progress=100,
                    message="转换完成",
                    result={
                        "file_id": req.file_id,
                        "download_url": f"/uploads/pdf2word_output/{req.file_id}.docx",
                        "output_filename": f"{req.file_id}.docx",
                        "page_count": page_count,
                        "ocr_pages": [],
                        "mode": "direct",
                    },
                )
                return

            update_task(task.task_id, progress=20, message="检测到扫描页或图片页，准备 OCR")
            page_texts, ocr_pages = await _ocr_pages(source_path, req.file_id, pages, task.task_id)
            update_task(task.task_id, progress=90, message="正在生成可编辑 Word 文档")
            await asyncio.to_thread(_build_text_docx, output_path, page_texts)
            update_task(
                task.task_id,
                status=TaskStatus.DONE,
                progress=100,
                message="转换完成（已自动识别并 AI 修复图片文字）",
                result={
                    "file_id": req.file_id,
                    "download_url": f"/uploads/pdf2word_output/{req.file_id}.docx",
                    "output_filename": f"{req.file_id}.docx",
                    "page_count": page_count,
                    "ocr_pages": ocr_pages,
                    "mode": "ocr_rebuilt",
                },
            )
        except Exception as e:
            if output_path.exists():
                output_path.unlink()
            update_task(
                task.task_id,
                status=TaskStatus.ERROR,
                progress=100,
                message=f"转换失败: {e}",
                result={},
            )

    bg.add_task(_run)
    return {"task_id": task.task_id}


@router.get("/task/{task_id}")
async def get_pdf2word_task(task_id: str):
    task = get_task(task_id)
    if not task:
        raise HTTPException(404, "任务不存在")
    return task
