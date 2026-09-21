"""
PPTX 生成服务 — 根据幻灯片结构生成专业格式的 .pptx 文件
"""
from pathlib import Path
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR


SLIDE_WIDTH = Inches(13.333)
SLIDE_HEIGHT = Inches(7.5)

FONT_TITLE = "SimHei"
FONT_BODY = "Microsoft YaHei"
FONT_CODE = "Consolas"

COLOR_TITLE = RGBColor(0x1A, 0x1A, 0x2E)
COLOR_BODY = RGBColor(0x33, 0x33, 0x33)
COLOR_ACCENT = RGBColor(0x00, 0x6B, 0x9F)
COLOR_LIGHT_BG = RGBColor(0xF5, 0xF5, 0xF5)


def build_pptx(slides: list[dict], output_path: Path) -> Path:
    output_path = Path(output_path)
    output_path.parent.mkdir(parents=True, exist_ok=True)

    prs = Presentation()
    prs.slide_width = SLIDE_WIDTH
    prs.slide_height = SLIDE_HEIGHT

    for slide_data in slides:
        layout = slide_data.get("layout", "content")
        if layout == "title":
            _add_title_slide(prs, slide_data)
        elif layout == "section":
            _add_section_slide(prs, slide_data)
        elif layout == "code":
            _add_code_slide(prs, slide_data)
        elif layout == "summary":
            _add_summary_slide(prs, slide_data)
        else:
            _add_content_slide(prs, slide_data)

        if slide_data.get("notes"):
            notes_slide = prs.slides[-1].notes_slide
            notes_slide.notes_text_frame.text = slide_data["notes"]

    prs.save(str(output_path))
    return output_path


def _set_font(run, name=FONT_BODY, size=Pt(24), color=COLOR_BODY, bold=False):
    run.font.name = name
    run.font.size = size
    run.font.color.rgb = color
    run.font.bold = bold


def _add_title_slide(prs, data):
    slide = prs.slides.add_slide(prs.slide_layouts[6])  # blank
    # Title
    left, top = Inches(1), Inches(2.2)
    txBox = slide.shapes.add_textbox(left, top, Inches(11), Inches(2))
    tf = txBox.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    run = p.add_run()
    run.text = data.get("title", "")
    _set_font(run, FONT_TITLE, Pt(44), COLOR_TITLE, bold=True)
    # Subtitle
    bullets = data.get("bullets", [])
    if bullets:
        p2 = tf.add_paragraph()
        p2.alignment = PP_ALIGN.CENTER
        p2.space_before = Pt(20)
        run2 = p2.add_run()
        run2.text = bullets[0]
        _set_font(run2, FONT_BODY, Pt(24), COLOR_BODY)


def _add_content_slide(prs, data):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    # Title bar
    left, top = Inches(0.8), Inches(0.5)
    txBox = slide.shapes.add_textbox(left, top, Inches(11), Inches(1))
    tf = txBox.text_frame
    p = tf.paragraphs[0]
    run = p.add_run()
    run.text = data.get("title", "")
    _set_font(run, FONT_TITLE, Pt(36), COLOR_ACCENT, bold=True)
    # Bullets
    bullets = data.get("bullets", [])
    if bullets:
        left, top = Inches(1.2), Inches(1.8)
        txBox = slide.shapes.add_textbox(left, top, Inches(10.5), Inches(5))
        tf = txBox.text_frame
        tf.word_wrap = True
        for i, bullet in enumerate(bullets):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            p.space_before = Pt(12)
            run = p.add_run()
            run.text = f"• {bullet}"
            _set_font(run, FONT_BODY, Pt(22), COLOR_BODY)


def _add_section_slide(prs, data):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    left, top = Inches(1), Inches(2.8)
    txBox = slide.shapes.add_textbox(left, top, Inches(11), Inches(1.5))
    tf = txBox.text_frame
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    run = p.add_run()
    run.text = data.get("title", "")
    _set_font(run, FONT_TITLE, Pt(40), COLOR_ACCENT, bold=True)


def _add_code_slide(prs, data):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    # Title
    left, top = Inches(0.8), Inches(0.4)
    txBox = slide.shapes.add_textbox(left, top, Inches(11), Inches(0.8))
    tf = txBox.text_frame
    p = tf.paragraphs[0]
    run = p.add_run()
    run.text = data.get("title", "")
    _set_font(run, FONT_TITLE, Pt(28), COLOR_ACCENT, bold=True)
    # Code block with background
    code_text = "\n".join(data.get("bullets", []))
    left, top = Inches(0.6), Inches(1.4)
    width, height = Inches(12), Inches(5.5)
    shape = slide.shapes.add_shape(1, left, top, width, height)  # rectangle
    shape.fill.solid()
    shape.fill.fore_color.rgb = COLOR_LIGHT_BG
    shape.line.fill.background()
    # Code text
    txBox = slide.shapes.add_textbox(Inches(0.9), Inches(1.6), Inches(11.4), Inches(5.2))
    tf = txBox.text_frame
    tf.word_wrap = True
    for i, line in enumerate(code_text.split("\n")):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        run = p.add_run()
        run.text = line
        _set_font(run, FONT_CODE, Pt(14), COLOR_BODY)


def _add_summary_slide(prs, data):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    # Title
    left, top = Inches(1), Inches(1.5)
    txBox = slide.shapes.add_textbox(left, top, Inches(11), Inches(1.2))
    tf = txBox.text_frame
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    run = p.add_run()
    run.text = data.get("title", "总结")
    _set_font(run, FONT_TITLE, Pt(36), COLOR_ACCENT, bold=True)
    # Summary bullets
    bullets = data.get("bullets", [])
    if bullets:
        left, top = Inches(1.5), Inches(3)
        txBox = slide.shapes.add_textbox(left, top, Inches(10), Inches(4))
        tf = txBox.text_frame
        tf.word_wrap = True
        for i, bullet in enumerate(bullets):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            p.space_before = Pt(14)
            p.alignment = PP_ALIGN.LEFT
            run = p.add_run()
            run.text = f"✓ {bullet}"
            _set_font(run, FONT_BODY, Pt(22), COLOR_BODY)
